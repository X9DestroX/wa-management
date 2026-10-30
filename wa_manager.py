import json, os, re, secrets, sqlite3, hashlib, threading, webbrowser, socket, time, shutil, ipaddress, urllib.request, urllib.error, base64, csv, io, unicodedata, math
from datetime import datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs, urlencode, quote

APP_NAME='W-A Systeme de Gestion'; DEFAULT_PORT=8765; DISCOVERY_PORT=8766; DISCOVERY_MAGIC='WA_MANAGER_DISCOVERY_V1'; DISCOVERY_INTERVAL=1.0; RUNTIME_ROLE='client'; RUNTIME_SERVER=''; RUNTIME_PORT=DEFAULT_PORT
DATA_DIR=os.path.join(os.path.expanduser('~'),'WA_Manager_System'); DB_PATH=os.path.join(DATA_DIR,'wa_manager.db'); CONFIG_PATH=os.path.join(DATA_DIR,'config.json'); BACKUP_DIR=os.path.join(DATA_DIR,'backups'); UI_PATH=os.path.join(os.path.dirname(os.path.abspath(__file__)),'ui.html')

os.makedirs(DATA_DIR,exist_ok=True); os.makedirs(BACKUP_DIR,exist_ok=True)
CLOUD_CONFIG_PATH=os.path.join(DATA_DIR,'cloud_config.json')

class CloudAPIError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status=status

def load_cloud_config():
    try:
        with open(CLOUD_CONFIG_PATH, encoding='utf-8-sig') as f:
            c=json.load(f)
        url=str(c.get('supabase_url','')).strip().rstrip('/')
        key=str(c.get('publishable_key','')).strip()
        enabled=bool(c.get('enabled')) and bool(url) and bool(key)
        if enabled:
            if not url.startswith('https://') or '.supabase.co' not in url:
                raise ValueError('L’URL Supabase doit commencer par https:// et se terminer par .supabase.co.')
            if key.startswith('sb_secret_') or key.startswith('sb_service_role_'):
                raise ValueError('Utilisez la clé Publishable/anon, jamais une clé Secret/service-role dans cette application.')
            # Legacy Supabase JWT-format keys can be distinguished by their payload role.
            if key.count('.') == 2:
                try:
                    part=key.split('.')[1]
                    part += '=' * ((4-len(part)%4)%4)
                    claims=json.loads(base64.urlsafe_b64decode(part.encode('ascii')).decode('utf-8'))
                    if claims.get('role') in ('service_role','supabase_admin'):
                        raise ValueError('Une clé service_role a été détectée. Supprimez-la de ce PC et utilisez la clé Publishable/anon.')
                except ValueError:
                    raise
                except Exception:
                    pass
        return {'enabled': enabled, 'supabase_url':url, 'publishable_key':key}
    except FileNotFoundError:
        return {'enabled':False}
    except Exception as e:
        return {'enabled':False,'error':str(e)}

def cloud_enabled():
    return bool(load_cloud_config().get('enabled'))

def _supabase_raw(cfg, path, method='GET', access_token=None, payload=None, prefer=None):
    url=cfg['supabase_url']+path
    headers={'apikey':cfg['publishable_key'],'Accept':'application/json','Content-Type':'application/json'}
    if access_token:
        headers['Authorization']='Bearer '+access_token
    if prefer:
        headers['Prefer']=prefer
    data=None if payload is None else json.dumps(payload,ensure_ascii=False).encode('utf-8')
    req=urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=25) as resp:
            raw=resp.read()
            if not raw: return None
            try: return json.loads(raw.decode('utf-8'))
            except Exception: return raw.decode('utf-8',errors='replace')
    except urllib.error.HTTPError as e:
        raw=e.read().decode('utf-8',errors='replace')
        try:
            j=json.loads(raw); msg=j.get('message') or j.get('msg') or j.get('error_description') or j.get('error') or raw
        except Exception: msg=raw or f'Erreur HTTP {e.code}'
        if e.code == 401:
            raise CloudAPIError(str(msg) or 'Session expirée.',401)
        if e.code == 403:
            raise CloudAPIError(str(msg) or 'Accès refusé.',403)
        if e.code == 404:
            raise CloudAPIError(str(msg) or 'Ressource introuvable.',404)
        if e.code == 409:
            raise CloudAPIError(str(msg) or 'Conflit de données.',409)
        raise CloudAPIError(str(msg)[:500], e.code if 400 <= e.code < 600 else 400)
    except urllib.error.URLError as e:
        raise CloudAPIError('Connexion Supabase impossible. Vérifiez Internet, l’URL du projet et votre connexion.',503)

def supabase_request(path, method='GET', session=None, payload=None, prefer=None):
    cfg=load_cloud_config()
    if not cfg.get('enabled'):
        raise CloudAPIError(cfg.get('error') or 'Configuration cloud absente. Exécutez configure_cloud.bat.',503)
    access=session.get('cloud_access_token') if session else None
    try:
        return _supabase_raw(cfg,path,method,access,payload,prefer)
    except CloudAPIError as e:
        if e.status != 401 or not session or not session.get('cloud_refresh_token') or path.startswith('/auth/v1/token?'):
            raise
        # Refresh the expired Supabase session and retry exactly once.
        refreshed=_supabase_raw(cfg,'/auth/v1/token?grant_type=refresh_token','POST',None,
                                {'refresh_token':session['cloud_refresh_token']})
        if not isinstance(refreshed,dict) or not refreshed.get('access_token'):
            raise CloudAPIError('Session expirée. Déconnectez-vous puis reconnectez-vous.',401)
        session['cloud_access_token']=refreshed['access_token']
        session['cloud_refresh_token']=refreshed.get('refresh_token',session['cloud_refresh_token'])
        return _supabase_raw(cfg,path,method,session['cloud_access_token'],payload,prefer)

def cloud_rest(table, params=None, method='GET', session=None, payload=None, prefer=None):
    query=('?'+urlencode(params,doseq=True)) if params else ''
    return supabase_request('/rest/v1/'+table+query,method,session,payload,prefer)

def iso_local(value):
    value=str(value or '').strip()
    if not value: return datetime.now().astimezone().isoformat(timespec='seconds')
    try:
        dt=datetime.fromisoformat(value.replace('Z','+00:00'))
        if dt.tzinfo is None: dt=dt.astimezone()
        return dt.isoformat(timespec='seconds')
    except Exception:
        return datetime.now().astimezone().isoformat(timespec='seconds')

def date_bounds(date_value):
    try: day=datetime.strptime(str(date_value),'%Y-%m-%d').date()
    except Exception: day=datetime.now().date()
    tz=datetime.now().astimezone().tzinfo
    start=datetime.combine(day,datetime.min.time()).replace(tzinfo=tz)
    end=datetime.combine(day+timedelta(days=1),datetime.min.time()).replace(tzinfo=tz)
    return start.isoformat(timespec='seconds'),end.isoformat(timespec='seconds')


def report_period(period, anchor):
    """Return inclusive start/end dates and an easy-to-display French label."""
    try: day=datetime.strptime(str(anchor), '%Y-%m-%d').date()
    except Exception: day=datetime.now().date()
    period='monthly' if str(period).lower() in ('monthly','month','mensuel','mensuelle') else 'weekly'
    if period=='monthly':
        start=day.replace(day=1)
        end=(start.replace(day=28)+timedelta(days=4)).replace(day=1)-timedelta(days=1)
        label='Rapport mensuel'
    else:
        start=day-timedelta(days=day.weekday())
        end=start+timedelta(days=6)
        label='Rapport hebdomadaire'
    start_at,_=date_bounds(start.isoformat())
    next_start,_=date_bounds((end+timedelta(days=1)).isoformat())
    return {'period':period,'start_date':start.isoformat(),'end_date':end.isoformat(),
            'start_at':start_at,'end_exclusive':next_start,
            'label':label, 'anchor_date':day.isoformat()}


def sum_field(rows, field):
    return round(sum(num(row.get(field)) for row in rows), 2)


def assemble_report(period_info, sales_rows, deposits_rows, expense_rows):
    sales_rows=list(sales_rows or []); deposits_rows=list(deposits_rows or []); expense_rows=list(expense_rows or [])
    gps=[x for x in sales_rows if str(x.get('device_type','')).upper()=='GPS']
    star=[x for x in sales_rows if str(x.get('device_type','')).upper()=='STARLINK']
    revenue=round(sum(num(x.get('total_amount') if x.get('total_amount') is not None else x.get('amount')) for x in sales_rows),2)
    subscriptions=sum_field(sales_rows,'subscription')
    sale_expenses=sum_field(sales_rows,'expenses')
    company_expenses=sum_field(expense_rows,'amount')
    total_expenses=round(sale_expenses+company_expenses,2)
    deposits=sum_field(deposits_rows,'amount')
    summary={
        'sales_count':len(sales_rows),'gps_count':len(gps),'starlink_count':len(star),
        'revenue':revenue,'subscriptions':subscriptions,'deposits':deposits,
        'sale_expenses':sale_expenses,'company_expenses':company_expenses,
        'total_expenses':total_expenses,'net_operating':round(revenue-total_expenses,2),
        'expenses_count':len(expense_rows),'deposits_count':len(deposits_rows),
    }
    return {**period_info,'summary':summary,'sales':sales_rows,'deposits':deposits_rows,'expenses':expense_rows}


def report_pdf_bytes(report):
    """Generate a multi-page landscape PDF; the PDF is created locally from authorized report data."""
    try:
        from io import BytesIO
        from xml.sax.saxutils import escape
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT, TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    except ImportError as e:
        raise CloudAPIError('Export PDF indisponible. Relancez build_exe.bat avec les dépendances de la version 1.7.3.', 500) from e

    buf=BytesIO()
    doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=12*mm,leftMargin=12*mm,topMargin=12*mm,bottomMargin=13*mm,title=str(report.get('label','Rapport')))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='WASection',parent=styles['Heading2'],fontName='Helvetica-Bold',fontSize=11,leading=14,textColor=colors.HexColor('#12355b'),spaceBefore=10,spaceAfter=5))
    styles.add(ParagraphStyle(name='WASmall',parent=styles['BodyText'],fontName='Helvetica',fontSize=7,leading=9,wordWrap='CJK'))
    styles.add(ParagraphStyle(name='WAHeader',parent=styles['WASmall'],fontName='Helvetica-Bold',textColor=colors.white))
    styles.add(ParagraphStyle(name='WATitle',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=19,leading=22,textColor=colors.HexColor('#12355b'),alignment=TA_LEFT,spaceAfter=4))
    styles.add(ParagraphStyle(name='WAMeta',parent=styles['BodyText'],fontName='Helvetica',fontSize=9,leading=12,textColor=colors.HexColor('#5d6878'),spaceAfter=8))

    def fmt(v): return f"{num(v):,.0f} FCFA".replace(',', ' ')
    def p(v, style='WASmall'):
        value='-' if v is None or v=='' else str(v)
        return Paragraph(escape(value).replace('\n','<br/>'),styles[style])
    def make_table(headers, rows, widths, font=7):
        data=[[Paragraph(escape(str(h)),styles['WAHeader']) for h in headers]]
        for row in rows:
            data.append([x if hasattr(x,'getPlainText') else p(x) for x in row])
        table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#12355b')),
            ('TEXTCOLOR',(0,0),(-1,0),colors.white),
            ('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#dce3ed')),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f6f8fb')]),
            ('VALIGN',(0,0),(-1,-1),'TOP'),
            ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
            ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
        ]))
        return table

    story=[]
    story.append(Paragraph('W-A Systeme de Gestion',styles['WATitle']))
    story.append(Paragraph(escape(str(report.get('label','Rapport')))+' — du '+escape(str(report.get('start_date','')))+' au '+escape(str(report.get('end_date',''))),styles['WAMeta']))
    story.append(Paragraph('Synthèse financière',styles['WASection']))
    sm=report['summary']
    summary_rows=[
        ['Ventes enregistrées',str(sm['sales_count']),'GPS',str(sm['gps_count']),'Starlink',str(sm['starlink_count'])],
        ['Chiffre d’affaires',fmt(sm['revenue']),'Abonnements',fmt(sm['subscriptions']),'Dépôts financiers',fmt(sm['deposits'])],
        ['Dépenses entreprise',fmt(sm['company_expenses']),'Dépenses des ventes',fmt(sm['sale_expenses']),'Total dépenses',fmt(sm['total_expenses'])],
        ['Position opérationnelle',fmt(sm['net_operating']),'Dépôts enregistrés',str(sm['deposits_count']),'Dépenses enregistrées',str(sm['expenses_count'])],
    ]
    story.append(make_table(['Indicateur','Valeur','Indicateur','Valeur','Indicateur','Valeur'],summary_rows,[33*mm,31*mm,30*mm,31*mm,32*mm,31*mm]))
    story.append(Paragraph('Détail des ventes',styles['WASection']))
    gps_rows=[]; star_rows=[]; other_rows=[]
    for x in report['sales']:
        typ=str(x.get('device_type') or '').upper()
        dt=str(x.get('sold_at') or '').replace('T',' ')[:16]
        total=x.get('total_amount') if x.get('total_amount') is not None else x.get('amount')
        if typ=='GPS':
            gps_rows.append([dt,x.get('client_name') or '-',x.get('mobile') or '-',x.get('sim_number') or x.get('device_phone') or '-',x.get('gps_type') or '-',fmt(x.get('amount')),fmt(x.get('subscription')),fmt(total),x.get('address') or '-',x.get('notes') or '-'])
        elif typ=='STARLINK':
            star_rows.append([dt,x.get('device_name') or 'Starlink',str(x.get('quantity') or 1),fmt(x.get('amount')),x.get('notes') or '-'])
        else:
            other_rows.append([dt,x.get('device_type') or '-',x.get('device_name') or '-',str(x.get('quantity') or 1),fmt(total),x.get('notes') or '-'])
    if gps_rows:
        story.append(Paragraph('Ventes GPS',styles['Heading3']))
        story.append(make_table(['Date et heure','Client','Téléphone','N° SIM','Type GPS','Montant','Abonnement','Total payé','Adresse','Note'],gps_rows,[22*mm,27*mm,20*mm,17*mm,13*mm,18*mm,18*mm,20*mm,36*mm,36*mm]))
    if star_rows:
        story.append(Paragraph('Ventes Starlink',styles['Heading3']))
        story.append(make_table(['Date et heure','Produit','Quantité','Montant à payer','Note'],star_rows,[35*mm,48*mm,20*mm,35*mm,102*mm]))
    if other_rows:
        story.append(Paragraph('Autres ventes',styles['Heading3']))
        story.append(make_table(['Date et heure','Catégorie','Produit','Quantité','Total payé','Note'],other_rows,[30*mm,30*mm,52*mm,18*mm,32*mm,78*mm]))
    if not report['sales']: story.append(Paragraph('Aucune vente enregistrée sur cette période.',styles['WAMeta']))

    story.append(Paragraph('Dépôts financiers',styles['WASection']))
    deposit_rows=[[str(x.get('created_at') or '').replace('T',' ')[:16],fmt(x.get('amount')),x.get('department') or 'Finance',x.get('reference') or '-',x.get('deposited_by') or '-',x.get('notes') or '-'] for x in report['deposits']]
    story.append(make_table(['Date','Montant','Service','Référence','Déposé par','Notes'],deposit_rows,[24*mm,25*mm,27*mm,30*mm,28*mm,76*mm]) if deposit_rows else Paragraph('Aucun dépôt enregistré sur cette période.',styles['WAMeta']))

    story.append(Paragraph('Dépenses de l’entreprise',styles['WASection']))
    expense_rows=[[str(x.get('incurred_at') or '').replace('T',' ')[:16],fmt(x.get('amount')),x.get('category') or '-',x.get('paid_to') or '-',x.get('payment_method') or '-',x.get('reference') or '-',x.get('description') or x.get('notes') or '-'] for x in report['expenses']]
    story.append(make_table(['Date','Montant','Catégorie','Bénéficiaire','Paiement','Référence','Description / notes'],expense_rows,[24*mm,25*mm,28*mm,29*mm,27*mm,29*mm,78*mm]) if expense_rows else Paragraph('Aucune dépense entreprise enregistrée sur cette période.',styles['WAMeta']))
    if sm['sale_expenses']:
        story.append(Paragraph('Dépenses associées aux ventes (anciens enregistrements)',styles['WASection']))
        legacy=[ [str(x.get('sold_at') or '').replace('T',' ')[:16],x.get('device_type') or '-',fmt(x.get('expenses')),x.get('client_name') or x.get('device_name') or '-',x.get('notes') or '-'] for x in report['sales'] if num(x.get('expenses'))]
        if legacy: story.append(make_table(['Date','Type','Dépense','Client / appareil','Note'],legacy,[30*mm,30*mm,30*mm,55*mm,155*mm]))
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont('Helvetica',7); canvas.setFillColor(colors.HexColor('#7b8492'))
        canvas.drawString(12*mm,7*mm,'Document généré par W-A Systeme de Gestion')
        canvas.drawRightString(landscape(A4)[0]-12*mm,7*mm,'Page %d' % document.page)
        canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return buf.getvalue()


def cloud_period_rows(table, column, period_info, session):
    """Fetch all records in a period in 1000-row pages to avoid silently truncating reports."""
    out=[]; offset=0; limit=1000
    date_filter='('+column+'.gte.'+period_info['start_at']+','+column+'.lt.'+period_info['end_exclusive']+')'
    while True:
        page=cloud_rest(table,{'select':'*','and':date_filter,'order':column+'.asc','limit':str(limit),'offset':str(offset)},session=session) or []
        out.extend(page)
        if len(page)<limit: break
        offset += limit
    return out


def cloud_report(period, anchor, session):
    info=report_period(period,anchor)
    sales=cloud_period_rows('sales','sold_at',info,session)
    deposits=cloud_period_rows('deposits','created_at',info,session)
    expenses=cloud_period_rows('expenses','incurred_at',info,session)
    return assemble_report(info,sales,deposits,expenses)

def cloud_http_error(handler, exc):
    status=getattr(exc,'status',400)
    msg=str(exc) or 'La requête cloud a échoué.'
    return handler.send_json({'error':msg},status)

def cloud_function(action, session, payload=None):
    data={'action':action}
    if payload: data.update(payload)
    return supabase_request('/functions/v1/admin-users','POST',session,data)

def cloud_auth_event(event, session):
    """Best-effort authentication event log; business CRUD is logged by DB triggers."""
    try:
        return supabase_request('/functions/v1/auth-event','POST',session,{'event':event})
    except Exception:
        # Login/logout should not fail just because an optional audit endpoint is not deployed.
        return None

def make_cloud_sale(d, old=None):
    source=dict(old or {})
    source.update(d)
    typ=str(source.get('device_type','GPS')).strip() or 'GPS'
    calc=calculate_sale(source)
    amount=num(source.get('amount'))
    sold_at=iso_local(source.get('sold_at'))
    return {
        'device_type':typ,
        'device_name':str(source.get('device_name') or ''),
        'product_id':int(source['product_id']) if str(source.get('product_id','')).strip().isdigit() else None,
        'quantity':max(1,int(num(source.get('quantity')) or 1)),
        'client_name':str(source.get('client_name') or '') if typ.upper()=='GPS' else '',
        'mobile':str(source.get('mobile') or '') if typ.upper()=='GPS' else '',
        'device_phone':str(source.get('device_phone') or ''),
        'sim_number':str(source.get('sim_number') or '') if typ.upper()=='GPS' else '',
        'gps_type':str(source.get('gps_type') or '') if typ.upper()=='GPS' else '',
        'address':str(source.get('address') or '') if typ.upper()=='GPS' else '',
        'setup_by':str(source.get('setup_by') or ''),
        'amount':amount,
        'subscription':calc['subscription'],
        'subscription_done':bool(source.get('subscription_done')) if isinstance(source.get('subscription_done'),bool) else str(source.get('subscription_done','')).lower() in ('true','yes','1','oui'),
        'expenses':0,
        'sold_at':sold_at,
        'notes':str(source.get('notes') or ''),
        'device_price':0,
        'setup_price':0,
        'sim_cost':calc['sim_cost'],
        'activation_cost':calc['activation_cost'],
        'benefit':0,
        'next_month_activation':0,
        'subscription_months':calc['subscription_months'] if typ.upper()=='GPS' else 0,
        'actual_subscription_monthly':0,
        'billed_subscription_monthly':calc['billed_subscription_monthly'],
        'total_amount':calc['total_amount'],
    }

def cloud_session_user(session, profile):
    return {'id':profile['id'],'username':profile['username'],'role':profile['role'],
            'full_name':profile.get('full_name') or profile['username'],
            'email':profile.get('email') or ''}

def product_payload(data, old=None):
    old=old or {}
    category=str(data.get('category',old.get('category','')) or '').strip()
    product_name=str(data.get('product_name',old.get('product_name','')) or '').strip()
    barcode_qr=str(data.get('barcode_qr',old.get('barcode_qr','')) or '').strip()
    notes=str(data.get('notes',old.get('notes','')) or '').strip()
    selling_price=num(data.get('selling_price',old.get('selling_price',0)))
    try:
        stock_quantity=int(data.get('stock_quantity',old.get('stock_quantity',0)) or 0)
    except (TypeError,ValueError):
        raise ValueError('La quantité en stock doit être un nombre entier.')
    is_active=data.get('is_active',old.get('is_active',True))
    if isinstance(is_active,str): is_active=is_active.strip().lower() not in ('false','0','non','off','')
    if not category: raise ValueError('La catégorie du produit est obligatoire.')
    if not product_name: raise ValueError('Le nom du produit est obligatoire.')
    if selling_price < 0: raise ValueError('Le prix de vente ne peut pas être négatif.')
    if stock_quantity < 0: raise ValueError('La quantité en stock ne peut pas être négative.')
    return {'category':category,'product_name':product_name,'barcode_qr':barcode_qr,
            'selling_price':round(selling_price,2),'stock_quantity':stock_quantity,
            'notes':notes,'is_active':bool(is_active)}


def inventory_actor_username(session):
    return str((session or {}).get('username') or (session or {}).get('full_name') or 'utilisateur')

def now(): return datetime.now().strftime('%Y-%m-%d %H:%M:%S')
def hash_password(p,salt=None): salt=salt or secrets.token_hex(16); return salt+':'+hashlib.pbkdf2_hmac('sha256',p.encode(),bytes.fromhex(salt),120000).hex()
def verify_password(p,stored):
    try: salt,h=stored.split(':',1); return secrets.compare_digest(hash_password(p,salt).split(':',1)[1],h)
    except: return False
def db(): c=sqlite3.connect(DB_PATH,timeout=30); c.row_factory=sqlite3.Row; return c
def init_db():
    c=db()
    c.executescript("""CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT UNIQUE NOT NULL,password TEXT NOT NULL,role TEXT NOT NULL,full_name TEXT,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS products(id INTEGER PRIMARY KEY AUTOINCREMENT,category TEXT NOT NULL,product_name TEXT NOT NULL,barcode_qr TEXT NOT NULL DEFAULT '',selling_price REAL NOT NULL DEFAULT 0,stock_quantity INTEGER NOT NULL DEFAULT 0,notes TEXT NOT NULL DEFAULT '',is_active INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS stock_movements(id INTEGER PRIMARY KEY AUTOINCREMENT,product_id INTEGER NOT NULL,movement_type TEXT NOT NULL,quantity_delta INTEGER NOT NULL,reference_table TEXT NOT NULL DEFAULT '',reference_id TEXT NOT NULL DEFAULT '',actor_id INTEGER,actor_username TEXT NOT NULL DEFAULT '',notes TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS expenses(id INTEGER PRIMARY KEY AUTOINCREMENT,amount REAL NOT NULL,category TEXT NOT NULL DEFAULT 'Autre',description TEXT NOT NULL DEFAULT '',paid_to TEXT NOT NULL DEFAULT '',payment_method TEXT NOT NULL DEFAULT '',reference TEXT NOT NULL DEFAULT '',notes TEXT NOT NULL DEFAULT '',incurred_at TEXT NOT NULL,recorded_by TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS deposits(id INTEGER PRIMARY KEY AUTOINCREMENT,amount REAL NOT NULL,department TEXT,reference TEXT,deposited_by TEXT,notes TEXT,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS sales(id INTEGER PRIMARY KEY AUTOINCREMENT,device_type TEXT NOT NULL,device_name TEXT,client_name TEXT,mobile TEXT,address TEXT,setup_by TEXT,amount REAL NOT NULL DEFAULT 0,subscription REAL NOT NULL DEFAULT 0,subscription_done INTEGER NOT NULL DEFAULT 0,expenses REAL NOT NULL DEFAULT 0,sold_at TEXT NOT NULL,notes TEXT);
    CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT,action TEXT,details TEXT,created_at TEXT NOT NULL);""")
    # Add v1.1 calculation fields to existing installations without destroying data.
    cols={r['name'] for r in c.execute('PRAGMA table_info(sales)').fetchall()}
    migrations={
        'device_price':'REAL NOT NULL DEFAULT 0',
        'setup_price':'REAL NOT NULL DEFAULT 0',
        'sim_cost':'REAL NOT NULL DEFAULT 0',
        'activation_cost':'REAL NOT NULL DEFAULT 0',
        'benefit':'REAL NOT NULL DEFAULT 0',
        'next_month_activation':'REAL NOT NULL DEFAULT 0',
        "device_phone":"TEXT NOT NULL DEFAULT ''",
        "sim_number":"TEXT NOT NULL DEFAULT ''",
        "gps_type":"TEXT NOT NULL DEFAULT ''",
        'subscription_months':'INTEGER NOT NULL DEFAULT 1',
        'actual_subscription_monthly':'REAL NOT NULL DEFAULT 2000',
        'billed_subscription_monthly':'REAL NOT NULL DEFAULT 3000',
        'total_amount':'REAL NOT NULL DEFAULT 0',
        'product_id':'INTEGER',
        'quantity':'INTEGER NOT NULL DEFAULT 1'
    }
    for name,definition in migrations.items():
        if name not in cols:
            c.execute(f'ALTER TABLE sales ADD COLUMN {name} {definition}')
    c.execute("CREATE UNIQUE INDEX IF NOT EXISTS products_barcode_qr_unique ON products(lower(trim(barcode_qr))) WHERE trim(coalesce(barcode_qr,'')) <> ''")
    c.execute('CREATE INDEX IF NOT EXISTS stock_movements_product_created_idx ON stock_movements(product_id, created_at DESC)')
    # Backfill total_amount for existing records without changing their base amount.
    c.execute("UPDATE sales SET total_amount=CASE WHEN UPPER(device_type)='GPS' THEN amount + COALESCE(sim_cost,3000) + COALESCE(subscription,0) ELSE amount END WHERE total_amount=0")
    if c.execute('SELECT COUNT(*) n FROM users').fetchone()['n']==0:
        c.execute('INSERT INTO users(username,password,role,full_name,created_at) VALUES(?,?,?,?,?)',('admin',hash_password('admin123'),'Administrateur','Administrateur',now()))
        c.execute('INSERT INTO users(username,password,role,full_name,created_at) VALUES(?,?,?,?,?)',('hr',hash_password('hr123'),'RH','Utilisateur RH',now()))
    # Migrate role labels from previous English versions to the French roles.
    c.execute("UPDATE users SET role='Administrateur' WHERE role='Manager' AND username='admin'")
    c.execute("UPDATE users SET role='Gestionnaire' WHERE role='Manager'")
    c.execute("UPDATE users SET role='RH' WHERE role='HR'")
    c.commit(); c.close()


def num(value):
    try: return float(value or 0)
    except (TypeError, ValueError): return 0.0


def calculate_sale(d):
    """GPS pricing: SIM + activation is 3,000 FCFA and already includes month 1.
    ``subscription_months`` is the number of ADDITIONAL months paid.
    Total subscription = additional months × 3,000 FCFA.
    Total to pay = base amount + 3,000 FCFA + total additional subscription.
    """
    typ=str(d.get('device_type','GPS')).strip() or 'GPS'
    months=max(0, int(num(d.get('subscription_months')) or 0))
    quantity=max(1,int(num(d.get('quantity')) or 1))
    base=num(d.get('amount'))
    sim_activation=3000.0*quantity if typ.upper()=='GPS' else num(d.get('sim_cost'))
    billed_monthly=3000.0 if typ.upper()=='GPS' else num(d.get('billed_subscription_monthly'))
    subscription=billed_monthly * months * quantity if typ.upper()=='GPS' else num(d.get('subscription'))
    total_amount=base + sim_activation + subscription if typ.upper()=='GPS' else base
    return {'device_price':0.0,'setup_price':0.0,'sim_cost':sim_activation,
            'activation_cost':sim_activation,'subscription':subscription,
            'subscription_months':months,'actual_subscription_monthly':0.0,
            'billed_subscription_monthly':billed_monthly,'benefit':0.0,
            'next_month_activation':0.0,'gross_total':base,'total_amount':total_amount}


def backup_database():
    """Create a consistent SQLite backup on the server at startup and daily."""
    if not os.path.exists(DB_PATH):
        return None
    stamp=datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
    dest=os.path.join(BACKUP_DIR, f'wa_manager_{stamp}.db')
    try:
        src=sqlite3.connect(DB_PATH, timeout=30)
        out=sqlite3.connect(dest)
        with out:
            src.backup(out)
        out.close(); src.close()
        # Keep the latest 14 backups.
        files=sorted([os.path.join(BACKUP_DIR,f) for f in os.listdir(BACKUP_DIR) if f.endswith('.db')], key=os.path.getmtime, reverse=True)
        for old in files[14:]:
            try: os.remove(old)
            except OSError: pass
        return dest
    except Exception:
        try:
            if os.path.exists(dest): os.remove(dest)
        except OSError: pass
        return None


def start_backup_scheduler():
    def loop():
        last_day=None
        while True:
            day=datetime.now().strftime('%Y-%m-%d')
            if day != last_day:
                backup_database(); last_day=day
            time.sleep(60)
    threading.Thread(target=loop,daemon=True).start()


def local_ip():
    # Resolve a usable private LAN address without requiring Internet access.
    try:
        s=socket.socket(socket.AF_INET,socket.SOCK_DGRAM); s.connect(('10.255.255.255',1)); ip=s.getsockname()[0]; s.close()
        if ip and not ip.startswith('127.'):
            return ip
    except Exception:
        pass
    try:
        candidates=socket.gethostbyname_ex(socket.gethostname())[2]
        for ip in candidates:
            if ip and not ip.startswith('127.'):
                return ip
    except Exception:
        pass
    return '127.0.0.1'


def discover_server(timeout=2.5):
    """Find a W-A Systeme de Gestion server on the local subnet using UDP broadcast."""
    sock=socket.socket(socket.AF_INET,socket.SOCK_DGRAM)
    sock.setsockopt(socket.SOL_SOCKET,socket.SO_BROADCAST,1)
    sock.settimeout(0.35)
    try:
        payload=json.dumps({'magic':DISCOVERY_MAGIC,'port':DEFAULT_PORT,'request':'discover'}).encode()
        sock.sendto(payload,('<broadcast>',DISCOVERY_PORT))
        deadline=time.time()+timeout
        candidates=[]
        while time.time()<deadline:
            try:
                raw,addr=sock.recvfrom(2048); data=json.loads(raw.decode())
                if data.get('magic')==DISCOVERY_MAGIC and data.get('role')=='server':
                    candidates.append((addr[0],int(data.get('port',DEFAULT_PORT))))
            except socket.timeout: continue
            except Exception: continue
        return candidates[0] if candidates else None
    finally: sock.close()


def discovery_responder(port):
    def loop():
        sock=socket.socket(socket.AF_INET,socket.SOCK_DGRAM)
        sock.setsockopt(socket.SOL_SOCKET,socket.SO_REUSEADDR,1)
        try: sock.setsockopt(socket.SOL_SOCKET,socket.SO_BROADCAST,1)
        except OSError: pass
        try: sock.bind(('',DISCOVERY_PORT))
        except OSError: return
        while True:
            try:
                raw,addr=sock.recvfrom(2048); data=json.loads(raw.decode())
                if data.get('magic')==DISCOVERY_MAGIC and data.get('request')=='discover':
                    response=json.dumps({'magic':DISCOVERY_MAGIC,'role':'server','port':port,'name':APP_NAME,'host':socket.gethostname()}).encode()
                    sock.sendto(response,(addr[0],DISCOVERY_PORT))
            except Exception:
                continue
    threading.Thread(target=loop,daemon=True).start()


def save_cfg(server,port,mode='auto'):
    with open(CONFIG_PATH,'w',encoding='utf-8') as f:
        json.dump({'server':server,'port':port,'server_mode': mode=='server','network_mode':mode},f,indent=2)

LAN_ACCESS_PATH=os.path.join(DATA_DIR,'lan_access.json')

def load_lan_access():
    default={'enabled':False,'max_clients':3,'allowed_ips':[]}
    try:
        with open(LAN_ACCESS_PATH,encoding='utf-8') as f:
            data=json.load(f)
        if not isinstance(data,dict): return default
        return {'enabled':bool(data.get('enabled',False)),'max_clients':max(1,min(3,int(data.get('max_clients',3) or 3))),'allowed_ips':[str(x).strip() for x in data.get('allowed_ips',[]) if str(x).strip()]}
    except Exception: return default

def client_ip_allowed(ip):
    access=load_lan_access()
    if not access['enabled']: return True
    if ip in ('127.0.0.1','::1','localhost'): return True
    return ip in access['allowed_ips']

# Product CSV/XLSX import helpers (v1.7.7)
def _import_header_key(value):
    raw=unicodedata.normalize('NFKD',str(value or ''))
    raw=''.join(ch for ch in raw if not unicodedata.combining(ch)).lower()
    return ' '.join(re.findall(r'[a-z0-9]+',raw))

_IMPORT_HEADER_ALIASES={}
for _field,_aliases in {
    'category':['category','categorie','catégorie','categorie produit'],
    'product_name':['product name','product_name','nom du produit','nom produit','produit'],
    'barcode_qr':['barcode qr','barcode_qr','code-barres / qr code','code barres qr code','code barre qr','code barres','code barre','barcode','qr code','ean','gtin'],
    'selling_price':['selling price','selling_price','prix de vente','prix vente','prix','prix fcfa'],
    'stock_quantity':['stock quantity','stock_quantity','quantite en stock','quantité en stock','quantite stock','stock','quantity','qty'],
    'notes':['notes','note','remarques','observation'],
}.items():
    for _alias in _aliases:
        _IMPORT_HEADER_ALIASES[_import_header_key(_alias)]=_field


def _import_number(value, field_label):
    if value is None or (isinstance(value,str) and not value.strip()):
        raise ValueError(f'{field_label} est obligatoire.')
    if isinstance(value,bool):
        raise ValueError(f'{field_label} doit être un nombre.')
    if isinstance(value,(int,float)):
        number=float(value)
    else:
        raw=str(value).strip().replace('\u00a0','').replace('\u202f','').replace(' ','')
        raw=re.sub(r'(?i)(fcfa|xaf|cfa)', '', raw)
        raw=raw.replace("'",'')
        if not raw:
            raise ValueError(f'{field_label} est obligatoire.')
        if ',' in raw and '.' in raw:
            if raw.rfind(',')>raw.rfind('.'):
                raw=raw.replace('.','').replace(',','.')
            else:
                raw=raw.replace(',','')
        elif ',' in raw:
            parts=raw.split(',')
            if len(parts)>2:
                raw=''.join(parts)
            elif len(parts)==2 and len(parts[-1])==3 and parts[-1].isdigit():
                raw=''.join(parts)
            else:
                raw=raw.replace(',','.')
        elif raw.count('.')>1:
            parts=raw.split('.')
            if all(x.isdigit() for x in parts) and len(parts[-1])==3:
                raw=''.join(parts)
            else:
                raise ValueError(f'{field_label} n’est pas un nombre valide.')
        elif '.' in raw and raw.rsplit('.',1)[-1].isdigit() and len(raw.rsplit('.',1)[-1])==3:
            before,after=raw.rsplit('.',1)
            if before.isdigit(): raw=before+after
        try:
            number=float(raw)
        except Exception:
            raise ValueError(f'{field_label} n’est pas un nombre valide.')
    if not math.isfinite(number):
        raise ValueError(f'{field_label} doit être un nombre fini.')
    return number


def _normalize_import_row(source, row_number):
    errors=[]
    category=str(source.get('category') or '').strip()
    product_name=str(source.get('product_name') or '').strip()
    barcode=str(source.get('barcode_qr') or '').strip()
    notes=str(source.get('notes') or '').strip()
    try:
        price=_import_number(source.get('selling_price'), 'Prix de vente')
        if price<0: errors.append('Le prix de vente ne peut pas être négatif.')
    except ValueError as exc:
        price=0.0; errors.append(str(exc))
    try:
        q=_import_number(source.get('stock_quantity'), 'Quantité en stock')
        if q<0: errors.append('La quantité ne peut pas être négative.')
        if not q.is_integer(): errors.append('La quantité en stock doit être un entier.')
        quantity=int(q) if q>=0 and q.is_integer() else 0
    except ValueError as exc:
        quantity=0; errors.append(str(exc))
    if not category: errors.append('La catégorie est obligatoire.')
    if not product_name: errors.append('Le nom du produit est obligatoire.')
    return {'row_number':int(row_number),'category':category,'product_name':product_name,
            'barcode_qr':barcode,'selling_price':round(price,2),'stock_quantity':quantity,
            'notes':notes,'errors':errors,'valid':not errors}


def parse_product_import_file(filename, encoded_content):
    filename=str(filename or 'produits.csv')
    ext=os.path.splitext(filename.lower())[1]
    if ext not in ('.csv','.xlsx'):
        raise ValueError('Format non pris en charge. Utilisez un fichier .xlsx ou .csv (pas .xls).')
    if not isinstance(encoded_content,str) or not encoded_content:
        raise ValueError('Le fichier importé est vide ou illisible.')
    if len(encoded_content)>11_200_000:
        raise ValueError('Le fichier dépasse 8 Mo. Divisez le catalogue en plusieurs fichiers.')
    try:
        raw=base64.b64decode(encoded_content,validate=True)
    except Exception:
        raise ValueError('Le contenu du fichier est invalide.')
    if len(raw)>8*1024*1024:
        raise ValueError('Le fichier dépasse 8 Mo. Divisez le catalogue en plusieurs fichiers.')
    if ext=='.xlsx':
        try:
            from openpyxl import load_workbook
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            worksheet=workbook.active
            matrix=[]
            for i,row in enumerate(worksheet.iter_rows(values_only=True),start=1):
                matrix.append(list(row))
                if i>1002:
                    workbook.close()
                    raise ValueError('Le fichier contient plus de 1 000 lignes de produits. Divisez-le en plusieurs fichiers.')
            workbook.close()
        except ImportError:
            raise ValueError('Import Excel indisponible. Relancez build_exe.bat avec les dépendances de la version 1.7.7.')
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('Impossible de lire ce fichier Excel. Enregistrez-le au format .xlsx et réessayez.') from exc
    else:
        try:
            text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: text=raw.decode('cp1252')
            except Exception: raise ValueError('Le CSV doit être enregistré en UTF-8 ou en Windows-1252.')
        sample='\n'.join(text.splitlines()[:10])
        try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
        except Exception:
            class _Dialect(csv.excel): delimiter=','
            dialect=_Dialect
        matrix=list(csv.reader(io.StringIO(text),dialect))
    if not matrix or not any(str(v or '').strip() for v in matrix[0]):
        raise ValueError('La première ligne doit contenir les en-têtes des colonnes.')
    headers={}
    for idx,value in enumerate(matrix[0]):
        key=_IMPORT_HEADER_ALIASES.get(_import_header_key(value))
        if key and key not in headers: headers[key]=idx
    missing=[label for key,label in [('category','Catégorie'),('product_name','Nom du produit'),('selling_price','Prix de vente'),('stock_quantity','Quantité en stock')] if key not in headers]
    if missing:
        raise ValueError('Colonnes obligatoires introuvables : '+', '.join(missing)+'. Utilisez le modèle fourni dans l’application.')
    rows=[]
    for line_no,values in enumerate(matrix[1:],start=2):
        if not any(str(v or '').strip() for v in values): continue
        source={}
        for field,idx in headers.items():
            source[field]=values[idx] if idx<len(values) else ''
        rows.append(_normalize_import_row(source,line_no))
        if len(rows)>1000:
            raise ValueError('Le fichier contient plus de 1 000 lignes de produits. Divisez-le en plusieurs fichiers.')
    if not rows:
        raise ValueError('Aucun produit trouvé après la ligne des en-têtes.')
    return rows


def mark_import_duplicates(rows, existing_products):
    rows=[dict(r,errors=list(r.get('errors') or [])) for r in rows]
    existing_codes={}
    for product in existing_products or []:
        code=str(product.get('barcode_qr') or '').strip().casefold()
        if code: existing_codes[code]=product
    seen={}
    for row in rows:
        code=str(row.get('barcode_qr') or '').strip().casefold()
        if code:
            if code in existing_codes:
                p=existing_codes[code]
                row['errors'].append(f"Code déjà utilisé par « {p.get('product_name') or 'un produit existant'} ».")
            elif code in seen:
                row['errors'].append(f"Code en doublon dans le fichier (ligne {seen[code]}).")
            else:
                seen[code]=row.get('row_number','?')
        row['valid']=not row['errors']
    return rows


def normalize_import_submission(items):
    rows=[]
    if not isinstance(items,list): raise ValueError('Liste de produits invalide.')
    if len(items)>1000: raise ValueError('Import limité à 1 000 produits par fichier.')
    for i,item in enumerate(items,start=1):
        if not isinstance(item,dict): item={}
        rowno=item.get('row_number',i+1)
        source={key:item.get(key,'') for key in ('category','product_name','barcode_qr','selling_price','stock_quantity','notes')}
        try: rowno=int(rowno)
        except Exception: rowno=i+1
        rows.append(_normalize_import_row(source,rowno))
    return rows


def existing_barcode_label(existing_products):
    return {str(p.get('barcode_qr') or '').strip().casefold():p for p in (existing_products or []) if str(p.get('barcode_qr') or '').strip()}


class API(BaseHTTPRequestHandler):
    sessions={}

    def products_for_import(self, session=None, connection=None):
        if cloud_enabled():
            all_rows=[]
            for offset in range(0,10000,1000):
                batch=cloud_rest('products',{'select':'id,category,product_name,barcode_qr','order':'id.asc','limit':'1000','offset':str(offset)},session=session) or []
                all_rows.extend(batch)
                if len(batch)<1000: break
            return all_rows
        if connection is None:
            connection=db()
            close_after=True
        else:
            close_after=False
        try:
            return [dict(r) for r in connection.execute('SELECT id,category,product_name,barcode_qr FROM products')]
        finally:
            if close_after: connection.close()

    def import_preview(self, d, session=None, connection=None):
        rows=parse_product_import_file(d.get('filename',''),d.get('file_base64',''))
        existing=self.products_for_import(session=session,connection=connection)
        rows=mark_import_duplicates(rows,existing)
        valid=sum(1 for r in rows if r['valid'])
        return {'items':rows,'summary':{'total':len(rows),'valid':valid,'skipped':len(rows)-valid}}

    def import_confirm_cloud(self,d,session):
        submitted=normalize_import_submission(d.get('items',[]))
        existing=self.products_for_import(session=session)
        rows=mark_import_duplicates(submitted,existing)
        imported=0; skipped=0; errors=[]
        for row in rows:
            if not row['valid']:
                skipped+=1; errors.append({'row_number':row['row_number'],'errors':row['errors']}); continue
            try:
                record=product_payload(row)
                cloud_rest('products',method='POST',session=session,payload=record,prefer='return=representation')
                imported+=1
                if row.get('barcode_qr'):
                    existing.append({'barcode_qr':row['barcode_qr'],'product_name':row['product_name']})
            except (CloudAPIError,ValueError) as exc:
                skipped+=1
                errors.append({'row_number':row['row_number'],'errors':[str(exc)[:240]]})
        return {'ok':True,'imported':imported,'skipped':skipped,'errors':errors}

    def cloud_get(self,p):
        if p=='/api/health':
            return self.send_json({'ok':True,'app':APP_NAME,'backend':'supabase'})
        u=self.auth()
        if not u: return self.send_json({'error':'Non autorisé'},401)
        try:
            if p=='/api/network':
                return self.send_json({'role':'cloud','server':'Supabase','port':443,'hostname':socket.gethostname(),
                                       'local_ip':local_ip(),'online':True,'access_enabled':True,'max_clients':0})
            if p=='/api/products':
                rows=cloud_rest('products',{'select':'*','order':'category.asc,product_name.asc','limit':'1000'},session=u) or []
                return self.send_json({'items':rows})
            if p=='/api/stock-movements':
                rows=cloud_rest('stock_movements',{'select':'*','order':'created_at.desc','limit':'300'},session=u) or []
                products=cloud_rest('products',{'select':'id,category,product_name','limit':'1000'},session=u) or []
                by_id={str(x.get('id')):x for x in products}
                for row in rows:
                    product=by_id.get(str(row.get('product_id')),{})
                    row['product_name']=product.get('product_name','Produit archivé')
                    row['category']=product.get('category','')
                return self.send_json({'items':rows})
            if p=='/api/sales':
                rows=cloud_rest('sales',{'select':'*','order':'id.desc','limit':'200'},session=u) or []
                return self.send_json({'items':rows})
            if p=='/api/deposits':
                rows=cloud_rest('deposits',{'select':'*','order':'id.desc','limit':'200'},session=u) or []
                return self.send_json({'items':rows})
            if p=='/api/expenses':
                rows=cloud_rest('expenses',{'select':'*','order':'incurred_at.desc','limit':'500'},session=u) or []
                return self.send_json({'items':rows})
            if p in ('/api/reports','/api/reports/pdf'):
                query=parse_qs(urlparse(self.path).query)
                period=query.get('period',['weekly'])[0]
                selected=query.get('date',[datetime.now().strftime('%Y-%m-%d')])[0]
                report=cloud_report(period,selected,u)
                if p.endswith('/pdf'):
                    return self.send_pdf(report)
                return self.send_json(report)
            if p=='/api/users':
                if u.get('role') in ('Administrateur','Superadmin'):
                    result=cloud_function('list',u)
                    return self.send_json({'items':result.get('users',[])})
                rows=cloud_rest('profiles',{'select':'id,username,role,full_name,created_at','order':'username.asc'},session=u) or []
                return self.send_json({'items':rows})
            if p=='/api/audit':
                if u.get('role') not in ('Administrateur','Superadmin'):
                    return self.send_json({'error':'La consultation des journaux est réservée à l’administrateur ou au Superadmin.'},403)
                fields='id,username,action,details,created_at,actor_id,actor_email,actor_role,entity_table,entity_id,operation,before_data,after_data'
                rows=cloud_rest('audit',{'select':fields,'order':'created_at.desc','limit':'500'},session=u) or []
                return self.send_json({'items':rows})
            if p=='/api/dashboard':
                query=parse_qs(urlparse(self.path).query)
                selected=query.get('date',[datetime.now().strftime('%Y-%m-%d')])[0]
                start,end=date_bounds(selected)
                sales_rows=cloud_rest('sales',{'select':'*','and':'(sold_at.gte.'+start+',sold_at.lt.'+end+')','order':'id.desc','limit':'200'},session=u) or []
                dep_rows=cloud_rest('deposits',{'select':'*','and':'(created_at.gte.'+start+',created_at.lt.'+end+')','order':'id.desc','limit':'200'},session=u) or []
                revenue=sum(num(x.get('total_amount') if x.get('total_amount') is not None else x.get('amount')) for x in sales_rows)
                subs=sum(num(x.get('subscription')) for x in sales_rows)
                sale_exp=sum(num(x.get('expenses')) for x in sales_rows)
                expense_rows=cloud_rest('expenses',{'select':'*','and':'(incurred_at.gte.'+start+',incurred_at.lt.'+end+')','order':'id.desc','limit':'200'},session=u) or []
                company_exp=sum(num(x.get('amount')) for x in expense_rows)
                exp=sale_exp+company_exp
                dep_total=sum(num(x.get('amount')) for x in dep_rows)
                gps=sum(1 for x in sales_rows if str(x.get('device_type','')).upper()=='GPS')
                star=sum(1 for x in sales_rows if str(x.get('device_type','')).upper()=='STARLINK')
                return self.send_json({'deposits':dep_total,'revenue':revenue,'subscriptions':subs,'expenses':exp,
                                       'sale_expenses':sale_exp,'company_expenses':company_exp,
                                       'gps':gps,'starlink':star,'recent':sales_rows[:8],'recent_deposits':dep_rows[:8]})
            return self.send_json({'error':'Introuvable'},404)
        except CloudAPIError as e:
            return cloud_http_error(self,e)

    def cloud_post(self,p,d):
        if p=='/api/login':
            cfg=load_cloud_config()
            if not cfg.get('enabled'):
                return self.send_json({'error':cfg.get('error') or 'Configuration cloud absente. Exécutez configure_cloud.bat.'},503)
            email=str(d.get('email') or d.get('username') or '').strip().lower()
            password=str(d.get('password') or '')
            if '@' not in email:
                return self.send_json({'error':'Saisissez votre adresse e-mail professionnelle et votre mot de passe.'},400)
            try:
                session_data=supabase_request('/auth/v1/token?grant_type=password','POST',None,
                                              {'email':email,'password':password})
                if not isinstance(session_data,dict) or not session_data.get('access_token') or not session_data.get('user'):
                    return self.send_json({'error':'Réponse de connexion Supabase invalide.'},502)
                cloud_session={'cloud_access_token':session_data['access_token'],
                               'cloud_refresh_token':session_data.get('refresh_token','')}
                uid=session_data['user'].get('id')
                rows=cloud_rest('profiles',{'select':'id,username,full_name,role,created_at','id':'eq.'+str(uid)},session=cloud_session) or []
                if not rows:
                    return self.send_json({'error':'Votre compte est authentifié mais aucun profil employé ne lui est associé. Contactez l’administrateur.'},403)
                profile=rows[0]
                if profile.get('role') not in ('RH','Gestionnaire','Administrateur','Superadmin'):
                    return self.send_json({'error':'Votre rôle d’employé est invalide. Contactez l’administrateur.'},403)
                local_token=secrets.token_urlsafe(32)
                cloud_session.update({'id':profile['id'],'username':profile['username'],'role':profile['role'],
                                      'full_name':profile.get('full_name') or profile['username'],
                                      'email':email,'created_at':profile.get('created_at','')})
                self.sessions[local_token]=cloud_session
                cloud_auth_event('login',cloud_session)
                return self.send_json({'token':local_token,'user':cloud_session_user(cloud_session,profile)})
            except CloudAPIError as e:
                if e.status in (400,401): return self.send_json({'error':'Adresse e-mail ou mot de passe incorrect.'},401)
                return cloud_http_error(self,e)

        u=self.auth()
        if not u: return self.send_json({'error':'Non autorisé'},401)
        if p=='/api/logout':
            cloud_auth_event('logout',u)
            self.sessions.pop(self.headers.get('Authorization','').replace('Bearer ',''),None)
            return self.send_json({'ok':True})
        try:
            if p=='/api/products/import-preview':
                preview=self.import_preview(d,session=u)
                return self.send_json(preview)
            if p=='/api/products/import-confirm':
                result=self.import_confirm_cloud(d,u)
                return self.send_json(result)
            if p=='/api/products':
                record=product_payload(d)
                created=cloud_rest('products',method='POST',session=u,payload=record,prefer='return=representation') or []
                item=created[0] if isinstance(created,list) and created else (created if isinstance(created,dict) else {})
                return self.send_json({'ok':True,'item':item})
            m=re.fullmatch(r'/api/products/([^/]+)/edit',p)
            if m:
                old_rows=cloud_rest('products',{'select':'*','id':'eq.'+m.group(1),'limit':'1'},session=u) or []
                if not old_rows: return self.send_json({'error':'Produit introuvable.'},404)
                record=product_payload(d,old_rows[0])
                updated=cloud_rest('products',{'id':'eq.'+m.group(1)},'PATCH',u,record,'return=representation') or []
                item=updated[0] if isinstance(updated,list) and updated else {}
                return self.send_json({'ok':True,'item':item})
            if p=='/api/deposits':
                record={'amount':num(d.get('amount')),'department':d.get('department') or 'Finance',
                        'reference':d.get('reference') or '',
                        'deposited_by':d.get('deposited_by') or u.get('full_name') or u.get('username',''),
                        'notes':d.get('notes') or '','created_at':iso_local(d.get('created_at'))}
                cloud_rest('deposits',method='POST',session=u,payload=record,prefer='return=representation')
                return self.send_json({'ok':True})
            m=re.fullmatch(r'/api/deposits/([^/]+)/edit',p)
            if m:
                record={'amount':num(d.get('amount')),'department':d.get('department') or 'Finance',
                        'reference':d.get('reference') or '','deposited_by':d.get('deposited_by') or '',
                        'notes':d.get('notes') or '','created_at':iso_local(d.get('created_at'))}
                cloud_rest('deposits',{'id':'eq.'+m.group(1)},'PATCH',u,record,'return=representation')
                return self.send_json({'ok':True})
            if p=='/api/expenses':
                amount=num(d.get('amount'))
                if amount < 0: return self.send_json({'error':'Le montant de la dépense ne peut pas être négatif.'},400)
                record={'amount':amount,'category':str(d.get('category') or 'Autre'),'description':str(d.get('description') or ''),
                        'paid_to':str(d.get('paid_to') or ''),'payment_method':str(d.get('payment_method') or ''),
                        'reference':str(d.get('reference') or ''),'notes':str(d.get('notes') or ''),
                        'incurred_at':iso_local(d.get('incurred_at')),'recorded_by':u.get('full_name') or u.get('username','')}
                cloud_rest('expenses',method='POST',session=u,payload=record,prefer='return=representation')
                return self.send_json({'ok':True})
            m=re.fullmatch(r'/api/expenses/([^/]+)/edit',p)
            if m:
                amount=num(d.get('amount'))
                if amount < 0: return self.send_json({'error':'Le montant de la dépense ne peut pas être négatif.'},400)
                record={'amount':amount,'category':str(d.get('category') or 'Autre'),'description':str(d.get('description') or ''),
                        'paid_to':str(d.get('paid_to') or ''),'payment_method':str(d.get('payment_method') or ''),
                        'reference':str(d.get('reference') or ''),'notes':str(d.get('notes') or ''),
                        'incurred_at':iso_local(d.get('incurred_at'))}
                cloud_rest('expenses',{'id':'eq.'+m.group(1)},'PATCH',u,record,'return=representation')
                return self.send_json({'ok':True})
            if p=='/api/sales':
                product_id=str(d.get('product_id') or '').strip()
                if not product_id.isdigit():
                    return self.send_json({'error':'Sélectionnez un produit du catalogue avant d’enregistrer la vente.'},400)
                qty=max(1,int(num(d.get('quantity')) or 1))
                products=cloud_rest('products',{'select':'id,category,product_name,selling_price,stock_quantity,is_active','id':'eq.'+product_id,'limit':'1'},session=u) or []
                if not products or not products[0].get('is_active',True): return self.send_json({'error':'Le produit sélectionné est introuvable ou archivé.'},404)
                product=products[0]
                if int(product.get('stock_quantity') or 0)<qty:
                    return self.send_json({'error':f"Stock insuffisant : {product.get('stock_quantity',0)} unité(s) disponible(s)."},409)
                d['product_id']=int(product['id']); d['quantity']=qty
                d['device_type']=product['category']; d['device_name']=product['product_name']
                d['amount']=num(product['selling_price'])*qty
                record=make_cloud_sale(d)
                created=cloud_rest('sales',method='POST',session=u,payload=record,prefer='return=representation')
                return self.send_json({'ok':True,'item':created[0] if isinstance(created,list) and created else None})
            m=re.fullmatch(r'/api/sales/([^/]+)/edit',p)
            if m:
                old_rows=cloud_rest('sales',{'select':'*','id':'eq.'+m.group(1),'limit':'1'},session=u) or []
                if not old_rows: return self.send_json({'error':'Vente introuvable.'},404)
                old_sale=old_rows[0]
                if 'product_id' in d and str(d.get('product_id') or '').strip():
                    product_id=str(d.get('product_id')).strip()
                    if not product_id.isdigit(): return self.send_json({'error':'Produit invalide.'},400)
                    qty=max(1,int(num(d.get('quantity')) or 1))
                    products=cloud_rest('products',{'select':'id,category,product_name,selling_price,stock_quantity,is_active','id':'eq.'+product_id,'limit':'1'},session=u) or []
                    if not products: return self.send_json({'error':'Le produit sélectionné est introuvable.'},404)
                    product=products[0]
                    same_product=str(old_sale.get('product_id') or '')==product_id
                    same_quantity=int(old_sale.get('quantity') or 1)==qty
                    if not product.get('is_active',True) and not (same_product and same_quantity):
                        return self.send_json({'error':'Ce produit est archivé et ne peut pas être utilisé pour modifier le produit ou la quantité de cette vente.'},409)
                    # Keep the original transaction price when only notes/date are edited; a later catalogue price change must not rewrite history.
                    if not (same_product and same_quantity):
                        available=int(product.get('stock_quantity') or 0)
                        if same_product: available += int(old_sale.get('quantity') or 1)
                        if available < qty: return self.send_json({'error':f'Stock insuffisant : {available} unité(s) disponible(s) pour cette modification.'},409)
                        d['amount']=num(product['selling_price'])*qty
                    d['product_id']=int(product['id']); d['quantity']=qty
                    d['device_type']=product['category']; d['device_name']=product['product_name']
                elif 'product_id' in d:
                    d['product_id']=None; d['quantity']=max(1,int(num(d.get('quantity')) or 1))
                record=make_cloud_sale(d,old_sale)
                updated=cloud_rest('sales',{'id':'eq.'+m.group(1)},'PATCH',u,record,'return=representation')
                return self.send_json({'ok':True,'item':updated[0] if isinstance(updated,list) and updated else None})
            if p=='/api/users':
                if u.get('role') not in ('Administrateur','Superadmin'): return self.send_json({'error':'Accès réservé à l’administrateur ou au Superadmin'},403)
                payload={'email':d.get('email',''),'username':d.get('username',''),
                         'full_name':d.get('full_name',''),'role':d.get('role','RH'),
                         'password':d.get('password','')}
                result=cloud_function('create',u,payload)
                return self.send_json({'ok':True,'result':result})
            m=re.fullmatch(r'/api/users/([^/]+)/edit',p)
            if m:
                if u.get('role') not in ('Administrateur','Superadmin'): return self.send_json({'error':'Accès réservé à l’administrateur ou au Superadmin'},403)
                payload={'id':m.group(1),'username':d.get('username',''),'full_name':d.get('full_name',''),
                         'email':d.get('email',''),'role':d.get('role','RH'),'password':d.get('password','')}
                result=cloud_function('update',u,payload)
                return self.send_json({'ok':True,'result':result})
            return self.send_json({'error':'Introuvable'},404)
        except CloudAPIError as e:
            return cloud_http_error(self,e)
        except Exception as e:
            return self.send_json({'error':'Erreur cloud : '+str(e)[:300]},400)

    def cloud_delete(self,p,u):
        if not u: return self.send_json({'error':'Non autorisé'},401)
        try:
            m=re.fullmatch(r'/api/products/([^/]+)',p)
            if m:
                updated=cloud_rest('products',{'id':'eq.'+m.group(1)},'PATCH',u,{'is_active':False},'return=representation') or []
                if not updated: return self.send_json({'error':'Produit introuvable.'},404)
                return self.send_json({'ok':True,'archived':True})
            m=re.fullmatch(r'/api/users/([^/]+)',p)
            if m:
                if u.get('role') not in ('Administrateur','Superadmin'): return self.send_json({'error':'Accès réservé à l’administrateur ou au Superadmin'},403)
                result=cloud_function('delete',u,{'id':m.group(1)})
                return self.send_json({'ok':True,'result':result})
            m=re.fullmatch(r'/api/sales/([^/]+)',p)
            if m:
                cloud_rest('sales',{'id':'eq.'+m.group(1)},'DELETE',u,None,'return=representation')
                return self.send_json({'ok':True})
            m=re.fullmatch(r'/api/deposits/([^/]+)',p)
            if m:
                cloud_rest('deposits',{'id':'eq.'+m.group(1)},'DELETE',u,None,'return=representation')
                return self.send_json({'ok':True})
            m=re.fullmatch(r'/api/expenses/([^/]+)',p)
            if m:
                cloud_rest('expenses',{'id':'eq.'+m.group(1)},'DELETE',u,None,'return=representation')
                return self.send_json({'ok':True})
            return self.send_json({'error':'Introuvable'},404)
        except CloudAPIError as e:
            return cloud_http_error(self,e)

    def send_pdf(self, report):
        """Save the PDF to the Windows Downloads folder and return its exact path.

        A browser-blob download can be silently blocked or handled differently by the
        embedded pywebview browser. Saving via the local Python process makes the
        output location deterministic and visible to the user.
        """
        try:
            raw = report_pdf_bytes(report)
            period = 'mensuel' if report.get('period') == 'monthly' else 'hebdomadaire'
            start_date = str(report.get('start_date', '') or '').replace('/', '-').replace('\\', '-')
            end_date = str(report.get('end_date', '') or '').replace('/', '-').replace('\\', '-')
            stamp = datetime.now().strftime('%H-%M-%S')

            downloads_dir = os.path.join(os.path.expanduser('~'), 'Downloads')
            if os.name == 'nt':
                # Respect a redirected Windows Downloads folder (for example OneDrive).
                try:
                    import winreg
                    with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                        r'Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders') as key:
                        configured, _ = winreg.QueryValueEx(key, '{374DE290-123F-4565-9164-39C4925E467B}')
                    if configured:
                        downloads_dir = os.path.expandvars(configured)
                except (OSError, ImportError):
                    pass

            os.makedirs(downloads_dir, exist_ok=True)
            filename = f"rapport_{period}_{start_date}_{end_date}_{stamp}.pdf"
            saved_path = os.path.abspath(os.path.join(downloads_dir, filename))
            with open(saved_path, 'wb') as pdf_file:
                pdf_file.write(raw)

            return self.send_json({
                'ok': True,
                'filename': filename,
                'saved_path': saved_path,
                'folder': os.path.abspath(downloads_dir),
                'message': 'Le rapport PDF a été enregistré dans Téléchargements.'
            })
        except CloudAPIError as exc:
            return self.send_json({'error': str(exc)}, getattr(exc, 'status', 500))
        except Exception as exc:
            return self.send_json({
                'error': "Impossible d’enregistrer le PDF dans Téléchargements : " + str(exc)[:240]
            }, 500)

    def log_message(self,*a): pass
    def cors(self): self.send_header('Access-Control-Allow-Origin','*'); self.send_header('Access-Control-Allow-Headers','Content-Type, Authorization'); self.send_header('Access-Control-Allow-Methods','GET,POST,DELETE,OPTIONS')
    def send_json(self,data,status=200):
        raw=json.dumps(data).encode(); self.send_response(status); self.send_header('Content-Type','application/json'); self.cors(); self.send_header('Content-Length',str(len(raw))); self.end_headers(); self.wfile.write(raw)
    def body(self): n=int(self.headers.get('Content-Length','0')); return json.loads(self.rfile.read(n) or b'{}')
    def auth(self): return self.sessions.get(self.headers.get('Authorization','').replace('Bearer ',''))
    def network_allowed(self):
        if client_ip_allowed(self.client_address[0]): return True
        self.send_json({'error':'Cet ordinateur n’est pas autorisé à accéder à W-A Systeme de Gestion.'},403); return False
    def do_OPTIONS(self):
        if not self.network_allowed(): return
        self.send_response(204); self.cors(); self.end_headers()
    def do_GET(self):
        if not self.network_allowed(): return
        p=urlparse(self.path).path
        if p in ('/','/ui.html'):
            raw=open(UI_PATH,'rb').read(); self.send_response(200); self.send_header('Content-Type','text/html; charset=utf-8'); self.send_header('Content-Length',str(len(raw))); self.end_headers(); self.wfile.write(raw); return
        if cloud_enabled():
            return self.cloud_get(p)
        if p=='/api/health': return self.send_json({'ok':True,'app':APP_NAME})
        u=self.auth()
        if not u:return self.send_json({'error':'Non autorisé'},401)
        c=db()
        try:
            if p=='/api/network':
                access=load_lan_access(); return self.send_json({'role':RUNTIME_ROLE,'server':RUNTIME_SERVER or '127.0.0.1','port':RUNTIME_PORT,'hostname':socket.gethostname(),'local_ip':local_ip(),'online':True,'access_enabled':access['enabled'],'max_clients':access['max_clients']})
            if p=='/api/dashboard':
                date=parse_qs(urlparse(self.path).query).get('date',[datetime.now().strftime('%Y-%m-%d')])[0]
                dep=c.execute('SELECT COALESCE(SUM(amount),0)v FROM deposits WHERE date(created_at)=?',(date,)).fetchone()['v']; rev=c.execute('SELECT COALESCE(SUM(total_amount),0)v FROM sales WHERE date(sold_at)=?',(date,)).fetchone()['v']; sub=c.execute('SELECT COALESCE(SUM(subscription),0)v FROM sales WHERE date(sold_at)=?',(date,)).fetchone()['v']; sale_exp=c.execute('SELECT COALESCE(SUM(expenses),0)v FROM sales WHERE date(sold_at)=?',(date,)).fetchone()['v']; company_exp=c.execute('SELECT COALESCE(SUM(amount),0)v FROM expenses WHERE date(incurred_at)=?',(date,)).fetchone()['v']; exp=sale_exp+company_exp; gps=c.execute("SELECT COUNT(*)v FROM sales WHERE device_type='GPS' AND date(sold_at)=?",(date,)).fetchone()['v']; star=c.execute("SELECT COUNT(*)v FROM sales WHERE device_type='Starlink' AND date(sold_at)=?",(date,)).fetchone()['v']; recent=[dict(x) for x in c.execute('SELECT * FROM sales WHERE date(sold_at)=? ORDER BY id DESC LIMIT 8',(date,))]; rd=[dict(x) for x in c.execute('SELECT * FROM deposits WHERE date(created_at)=? ORDER BY id DESC LIMIT 8',(date,))]; return self.send_json({'deposits':dep,'revenue':rev,'subscriptions':sub,'expenses':exp,'sale_expenses':sale_exp,'company_expenses':company_exp,'gps':gps,'starlink':star,'recent':recent,'recent_deposits':rd})
            if p=='/api/products':return self.send_json({'items':[dict(x) for x in c.execute('SELECT * FROM products ORDER BY category,product_name')]})
            if p=='/api/stock-movements':
                rows=[dict(x) for x in c.execute('SELECT m.*,p.product_name,p.category FROM stock_movements m LEFT JOIN products p ON p.id=m.product_id ORDER BY m.created_at DESC,m.id DESC LIMIT 300')]
                return self.send_json({'items':rows})
            if p=='/api/deposits':return self.send_json({'items':[dict(x) for x in c.execute('SELECT * FROM deposits ORDER BY id DESC LIMIT 200')]})
            if p=='/api/expenses':return self.send_json({'items':[dict(x) for x in c.execute('SELECT * FROM expenses ORDER BY incurred_at DESC LIMIT 500')]})
            if p in ('/api/reports','/api/reports/pdf'):
                query=parse_qs(urlparse(self.path).query); period=query.get('period',['weekly'])[0]; selected=query.get('date',[datetime.now().strftime('%Y-%m-%d')])[0]
                info=report_period(period,selected); sales_rows=[dict(x) for x in c.execute('SELECT * FROM sales WHERE date(sold_at)>=? AND date(sold_at)<=? ORDER BY sold_at',(info['start_date'],info['end_date']))]
                dep_rows=[dict(x) for x in c.execute('SELECT * FROM deposits WHERE date(created_at)>=? AND date(created_at)<=? ORDER BY created_at',(info['start_date'],info['end_date']))]
                expense_rows=[dict(x) for x in c.execute('SELECT * FROM expenses WHERE date(incurred_at)>=? AND date(incurred_at)<=? ORDER BY incurred_at',(info['start_date'],info['end_date']))]
                report=assemble_report(info,sales_rows,dep_rows,expense_rows)
                if p.endswith('/pdf'):return self.send_pdf(report)
                return self.send_json(report)
            if p=='/api/sales':return self.send_json({'items':[dict(x) for x in c.execute('SELECT * FROM sales ORDER BY id DESC LIMIT 200')]})
            if p=='/api/users':
                return self.send_json({'items':[dict(x) for x in c.execute('SELECT id,username,role,full_name,created_at FROM users ORDER BY id')]})
            if p=='/api/audit':
                if u.get('role') not in ('Administrateur','Superadmin'):
                    return self.send_json({'error':'La consultation des journaux est réservée à l’administrateur ou au Superadmin.'},403)
                return self.send_json({'items':[dict(x) for x in c.execute('SELECT id,username,action,details,created_at FROM audit ORDER BY id DESC LIMIT 500')]})
            return self.send_json({'error':'Introuvable'},404)
        finally:c.close()
    def do_POST(self):
        if not self.network_allowed(): return
        p=urlparse(self.path).path; d=self.body()
        if cloud_enabled():
            return self.cloud_post(p,d)
        if p=='/api/login':
            c=db(); u=c.execute('SELECT * FROM users WHERE username=?',(d.get('username',''),)).fetchone()
            if not u or not verify_password(d.get('password',''),u['password']):
                c.close(); return self.send_json({'error':'Nom d’utilisateur ou mot de passe incorrect'},401)
            c.execute('INSERT INTO audit(username,action,details,created_at) VALUES(?,?,?,?)',
                      (u['username'],'Connexion réussie',json.dumps({'event':'login','role':u['role']},ensure_ascii=False),now()))
            c.commit(); tok=secrets.token_urlsafe(32)
            self.sessions[tok]={'id':u['id'],'username':u['username'],'role':u['role'],'full_name':u['full_name']}
            c.close(); return self.send_json({'token':tok,'user':self.sessions[tok]})
        u=self.auth()
        if not u:return self.send_json({'error':'Non autorisé'},401)
        if p=='/api/logout':
            c=db()
            c.execute('INSERT INTO audit(username,action,details,created_at) VALUES(?,?,?,?)',
                      (u.get('username',''),'Déconnexion',json.dumps({'event':'logout','role':u.get('role','')},ensure_ascii=False),now()))
            c.commit(); c.close()
            self.sessions.pop(self.headers.get('Authorization','').replace('Bearer ',''),None)
            return self.send_json({'ok':True})
        c=db(); cur=c.cursor()
        try:
            if p=='/api/products/import-preview':
                preview=self.import_preview(d,connection=c)
                return self.send_json(preview)
            if p=='/api/products/import-confirm':
                submitted=normalize_import_submission(d.get('items',[]))
                existing=self.products_for_import(connection=c)
                rows=mark_import_duplicates(submitted,existing)
                imported=0; skipped=0; errors=[]; stamp=now()
                for row in rows:
                    if not row['valid']:
                        skipped+=1; errors.append({'row_number':row['row_number'],'errors':row['errors']}); continue
                    try:
                        rec=product_payload(row)
                        cur.execute('INSERT INTO products(category,product_name,barcode_qr,selling_price,stock_quantity,notes,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)',
                            (rec['category'],rec['product_name'],rec['barcode_qr'],rec['selling_price'],rec['stock_quantity'],rec['notes'],1,stamp,stamp))
                        pid=cur.lastrowid
                        if rec['stock_quantity']:
                            cur.execute('INSERT INTO stock_movements(product_id,movement_type,quantity_delta,actor_id,actor_username,notes,created_at) VALUES(?,?,?,?,?,?,?)',
                                (pid,'stock_initial',rec['stock_quantity'],u.get('id'),u.get('username',''),'Stock initial importé par fichier',stamp))
                        imported+=1
                        if rec['barcode_qr']:
                            existing.append({'barcode_qr':rec['barcode_qr'],'product_name':rec['product_name']})
                    except (ValueError,sqlite3.IntegrityError) as exc:
                        skipped+=1; errors.append({'row_number':row['row_number'],'errors':[str(exc)[:240]]})
                c.commit()
                return self.send_json({'ok':True,'imported':imported,'skipped':skipped,'errors':errors})
            if p=='/api/products':
                rec=product_payload(d)
                stamp=now()
                cur.execute('INSERT INTO products(category,product_name,barcode_qr,selling_price,stock_quantity,notes,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)',
                    (rec['category'],rec['product_name'],rec['barcode_qr'],rec['selling_price'],rec['stock_quantity'],rec['notes'],1 if rec['is_active'] else 0,stamp,stamp))
                pid=cur.lastrowid
                if rec['stock_quantity']:
                    cur.execute('INSERT INTO stock_movements(product_id,movement_type,quantity_delta,actor_id,actor_username,notes,created_at) VALUES(?,?,?,?,?,?,?)',
                        (pid,'stock_initial',rec['stock_quantity'],u.get('id'),u.get('username',''),'Stock initial à la création du produit',stamp))
            elif re.fullmatch(r'/api/products/([^/]+)/edit',p):
                uid=re.fullmatch(r'/api/products/([^/]+)/edit',p).group(1)
                target=cur.execute('SELECT * FROM products WHERE id=?',(uid,)).fetchone()
                if not target:return self.send_json({'error':'Produit introuvable.'},404)
                rec=product_payload(d,dict(target)); old_qty=int(target['stock_quantity'])
                cur.execute('UPDATE products SET category=?,product_name=?,barcode_qr=?,selling_price=?,stock_quantity=?,notes=?,is_active=?,updated_at=? WHERE id=?',
                    (rec['category'],rec['product_name'],rec['barcode_qr'],rec['selling_price'],rec['stock_quantity'],rec['notes'],1 if rec['is_active'] else 0,now(),uid))
                delta=rec['stock_quantity']-old_qty
                if delta:
                    cur.execute('INSERT INTO stock_movements(product_id,movement_type,quantity_delta,actor_id,actor_username,notes,created_at) VALUES(?,?,?,?,?,?,?)',
                        (uid,'ajustement',delta,u.get('id'),u.get('username',''),'Ajustement manuel du stock',now()))
            elif p=='/api/deposits':cur.execute('INSERT INTO deposits(amount,department,reference,deposited_by,notes,created_at)VALUES(?,?,?,?,?,?)',(float(d.get('amount',0)),d.get('department','Finance'),d.get('reference',''),d.get('deposited_by','') or u['full_name'] or u['username'],d.get('notes',''),d.get('created_at') or now()))
            elif p.startswith('/api/deposits/') and p.endswith('/edit'):
                uid=p.split('/')[-2]
                if not uid.isdigit(): return self.send_json({'error':'Identifiant du dépôt invalide'},400)
                target=cur.execute('SELECT * FROM deposits WHERE id=?',(int(uid),)).fetchone()
                if not target:return self.send_json({'error':'Dépôt introuvable'},404)
                cur.execute('UPDATE deposits SET amount=?,department=?,reference=?,deposited_by=?,notes=?,created_at=? WHERE id=?',(
                    float(d.get('amount',target['amount'])),d.get('department',target['department'] or 'Finance'),d.get('reference',target['reference'] or ''),
                    d.get('deposited_by',target['deposited_by'] or ''),d.get('notes',target['notes'] or ''),d.get('created_at',target['created_at']),int(uid)))
            elif p=='/api/expenses':
                amount=num(d.get('amount'))
                if amount<0:return self.send_json({'error':'Le montant de la dépense ne peut pas être négatif.'},400)
                cur.execute('INSERT INTO expenses(amount,category,description,paid_to,payment_method,reference,notes,incurred_at,recorded_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',
                    (amount,d.get('category') or 'Autre',d.get('description') or '',d.get('paid_to') or '',d.get('payment_method') or '',d.get('reference') or '',d.get('notes') or '',d.get('incurred_at') or now(),u.get('full_name') or u.get('username',''),now()))
            elif re.fullmatch(r'/api/expenses/([^/]+)/edit',p):
                uid=re.fullmatch(r'/api/expenses/([^/]+)/edit',p).group(1)
                target=cur.execute('SELECT * FROM expenses WHERE id=?',(uid,)).fetchone()
                if not target:return self.send_json({'error':'Dépense introuvable'},404)
                amount=num(d.get('amount',target['amount']))
                if amount<0:return self.send_json({'error':'Le montant de la dépense ne peut pas être négatif.'},400)
                cur.execute('UPDATE expenses SET amount=?,category=?,description=?,paid_to=?,payment_method=?,reference=?,notes=?,incurred_at=? WHERE id=?',
                    (amount,d.get('category',target['category']),d.get('description',target['description']),d.get('paid_to',target['paid_to']),d.get('payment_method',target['payment_method']),d.get('reference',target['reference']),d.get('notes',target['notes']),d.get('incurred_at',target['incurred_at']),uid))
            elif p=='/api/sales':
                product_id=str(d.get('product_id') or '').strip()
                if not product_id.isdigit(): return self.send_json({'error':'Sélectionnez un produit du catalogue avant d’enregistrer la vente.'},400)
                qty=max(1,int(num(d.get('quantity')) or 1))
                product=cur.execute('SELECT * FROM products WHERE id=? AND is_active=1',(int(product_id),)).fetchone()
                if not product:return self.send_json({'error':'Le produit sélectionné est introuvable ou archivé.'},404)
                if int(product['stock_quantity'])<qty:return self.send_json({'error':f"Stock insuffisant : {product['stock_quantity']} unité(s) disponible(s)."},409)
                d['quantity']=qty; d['device_type']=product['category']; d['device_name']=product['product_name']
                d['amount']=num(product['selling_price'])*qty
                calc=calculate_sale(d); amount=calc['total_amount']
                cur.execute('''INSERT INTO sales(device_type,device_name,client_name,mobile,device_phone,sim_number,gps_type,address,setup_by,amount,subscription,subscription_done,expenses,sold_at,notes,device_price,setup_price,sim_cost,activation_cost,benefit,next_month_activation,subscription_months,actual_subscription_monthly,billed_subscription_monthly,total_amount,product_id,quantity)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                    (d.get('device_type'),d.get('device_name',''),d.get('client_name',''),d.get('mobile',''),d.get('sim_number',''),d.get('sim_number',''),d.get('gps_type',''),d.get('address',''),d.get('setup_by',''),
                     num(d.get('amount')),calc['subscription'],1 if d.get('subscription_done') else 0,0,d.get('sold_at') or now(),d.get('notes',''),
                     0,0,calc['sim_cost'],calc['activation_cost'],0,0,calc['subscription_months'],0,calc['billed_subscription_monthly'],calc['total_amount'],int(product_id),qty))
                sale_id=cur.lastrowid
                cur.execute('UPDATE products SET stock_quantity=stock_quantity-?,updated_at=? WHERE id=?',(qty,now(),int(product_id)))
                cur.execute('INSERT INTO stock_movements(product_id,movement_type,quantity_delta,reference_table,reference_id,actor_id,actor_username,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?)',
                    (int(product_id),'vente',-qty,'sales',str(sale_id),u.get('id'),u.get('username',''),'Sortie de stock liée à une vente',now()))
            elif p.startswith('/api/sales/') and p.endswith('/edit'):
                uid=p.split('/')[-2]
                if not uid.isdigit(): return self.send_json({'error':'Identifiant de vente invalide'},400)
                target=cur.execute('SELECT * FROM sales WHERE id=?',(int(uid),)).fetchone()
                if not target:return self.send_json({'error':'Vente introuvable'},404)
                # Preserve hidden pricing details; restore old stock and deduct the chosen product atomically in this SQLite transaction.
                calc_input=dict(target); calc_input.update(d)
                target_product=int(target['product_id']) if target['product_id'] is not None else None
                new_product=target_product
                if 'product_id' in d:
                    new_product=int(d['product_id']) if str(d.get('product_id') or '').isdigit() else None
                qty=max(1,int(num(d.get('quantity',target['quantity'] or 1)) or 1))
                if new_product is not None:
                    prod=cur.execute('SELECT * FROM products WHERE id=?',(new_product,)).fetchone()
                    if not prod:return self.send_json({'error':'Le produit sélectionné est introuvable.'},404)
                    same_product=target_product==new_product; same_quantity=int(target['quantity'] or 1)==qty
                    if not prod['is_active'] and not (same_product and same_quantity):return self.send_json({'error':'Ce produit est archivé et ne peut pas être utilisé pour modifier le produit ou la quantité de cette vente.'},409)
                    if not (same_product and same_quantity):
                        available=int(prod['stock_quantity']) + (int(target['quantity'] or 1) if same_product else 0)
                        if available<qty:return self.send_json({'error':f'Stock insuffisant : {available} unité(s) disponible(s) pour cette modification.'},409)
                        calc_input['amount']=num(prod['selling_price'])*qty
                    calc_input['device_type']=prod['category'];calc_input['device_name']=prod['product_name']
                calc_input['quantity']=qty
                calc=calculate_sale(calc_input)
                final_amount=num(calc_input.get('amount',target['amount']))
                if target_product is not None and (target_product!=new_product or int(target['quantity'] or 1)!=qty):
                    cur.execute('UPDATE products SET stock_quantity=stock_quantity+?,updated_at=? WHERE id=?',(int(target['quantity'] or 1),now(),target_product))
                    cur.execute('INSERT INTO stock_movements(product_id,movement_type,quantity_delta,reference_table,reference_id,actor_id,actor_username,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(target_product,'retour_vente',int(target['quantity'] or 1),'sales',str(uid),u.get('id'),u.get('username',''),'Restitution avant modification de vente',now()))
                if new_product is not None and (target_product!=new_product or int(target['quantity'] or 1)!=qty):
                    cur.execute('UPDATE products SET stock_quantity=stock_quantity-?,updated_at=? WHERE id=?',(qty,now(),new_product))
                    cur.execute('INSERT INTO stock_movements(product_id,movement_type,quantity_delta,reference_table,reference_id,actor_id,actor_username,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(new_product,'vente',-qty,'sales',str(uid),u.get('id'),u.get('username',''),'Sortie après modification de vente',now()))
                cur.execute('''UPDATE sales SET device_type=?,device_name=?,client_name=?,mobile=?,device_phone=?,sim_number=?,gps_type=?,address=?,setup_by=?,amount=?,subscription=?,subscription_done=?,expenses=?,sold_at=?,notes=?,device_price=?,setup_price=?,sim_cost=?,activation_cost=?,benefit=?,next_month_activation=?,subscription_months=?,actual_subscription_monthly=?,billed_subscription_monthly=?,total_amount=?,product_id=?,quantity=? WHERE id=?''',(
                    calc_input.get('device_type',target['device_type']),calc_input.get('device_name',target['device_name'] or ''),d.get('client_name',target['client_name'] or ''),
                    d.get('mobile',target['mobile'] or ''),d.get('device_phone',target['device_phone'] or ''),d.get('sim_number',target['sim_number'] if 'sim_number' in target.keys() else ''),d.get('gps_type',target['gps_type'] if 'gps_type' in target.keys() else ''),d.get('address',target['address'] or ''),
                    d.get('setup_by',target['setup_by'] or ''),final_amount,calc['subscription'],1 if d.get('subscription_done') else 0,
                    0,d.get('sold_at') or target['sold_at'],d.get('notes',target['notes'] or ''),0,0,
                    calc['sim_cost'],calc['activation_cost'],0,0,calc['subscription_months'],0,
                    calc['billed_subscription_monthly'],calc['total_amount'],new_product,qty,int(uid)))
            elif p=='/api/users':
                if u['role'] not in ('Administrateur','Superadmin'):return self.send_json({'error':'Accès réservé à l’administrateur ou au Superadmin'},403)
                role=str(d.get('role','')).strip()
                if role not in ('RH','Gestionnaire','Administrateur'):return self.send_json({'error':'Rôle invalide'},400)
                if not str(d.get('username','')).strip() or not str(d.get('password','')):return self.send_json({'error':'Nom d’utilisateur et mot de passe obligatoires'},400)
                cur.execute('INSERT INTO users(username,password,role,full_name,created_at)VALUES(?,?,?,?,?)',(d['username'].strip(),hash_password(d['password']),role,d.get('full_name','').strip(),now()))
            elif p.startswith('/api/users/') and p.endswith('/edit'):
                if u['role'] not in ('Administrateur','Superadmin'):
                    return self.send_json({'error':'Accès réservé à l’administrateur ou au Superadmin'},403)
                uid=p.split('/')[-2]
                if not uid.isdigit(): return self.send_json({'error':'Identifiant utilisateur invalide'},400)
                target=cur.execute('SELECT * FROM users WHERE id=?',(int(uid),)).fetchone()
                if not target:return self.send_json({'error':'Utilisateur introuvable'},404)
                username=str(d.get('username',target['username'])).strip()
                full_name=str(d.get('full_name',target['full_name'] or '')).strip()
                role=str(d.get('role',target['role'])).strip()
                if not username:return self.send_json({'error':'Le nom d’utilisateur est obligatoire'},400)
                if role not in ('RH','Gestionnaire','Administrateur'):return self.send_json({'error':'Rôle invalide'},400)
                if d.get('password',''):
                    cur.execute('UPDATE users SET username=?,password=?,role=?,full_name=? WHERE id=?',(username,hash_password(d['password']),role,full_name,int(uid)))
                else:
                    cur.execute('UPDATE users SET username=?,role=?,full_name=? WHERE id=?',(username,role,full_name,int(uid)))
            else:return self.send_json({'error':'Introuvable'},404)
            # Never store passwords in the audit log. Record fields/action only.
            audit_payload=dict(d)
            password_changed=bool(audit_payload.pop('password', '') or '')
            if p=='/api/sales': audit_action='Ajout : vente'
            elif p.startswith('/api/sales/') and p.endswith('/edit'): audit_action='Modification : vente'
            elif p=='/api/deposits': audit_action='Ajout : dépôt'
            elif p.startswith('/api/deposits/') and p.endswith('/edit'): audit_action='Modification : dépôt'
            elif p=='/api/expenses': audit_action='Ajout : dépense'
            elif p.startswith('/api/expenses/') and p.endswith('/edit'): audit_action='Modification : dépense'
            elif p=='/api/users': audit_action='Création d’un utilisateur'
            elif p.startswith('/api/users/') and p.endswith('/edit'): audit_action='Modification d’un utilisateur'
            else: audit_action=p
            audit_details={'endpoint':p,'fields':audit_payload}
            if password_changed: audit_details['password_changed']=True
            cur.execute('INSERT INTO audit(username,action,details,created_at)VALUES(?,?,?,?)',(u['username'],audit_action,json.dumps(audit_details,ensure_ascii=False),now())); c.commit(); return self.send_json({'ok':True})
        except sqlite3.IntegrityError:
            if p.startswith('/api/products'):
                return self.send_json({'error':'Ce code-barres / QR est déjà associé à un produit.'},409)
            return self.send_json({'error':'Ce nom d’utilisateur existe déjà'},400)
        except Exception as e:c.rollback(); return self.send_json({'error':str(e)},400)
        finally:c.close()
    def do_DELETE(self):
        if not self.network_allowed(): return
        p=urlparse(self.path).path; u=self.auth()
        if cloud_enabled():
            return self.cloud_delete(p,u)
        if not u:return self.send_json({'error':'Non autorisé'},401)
        c=db(); cur=c.cursor()
        try:
            product_match=re.fullmatch(r'/api/products/([^/]+)',p)
            if product_match:
                target=cur.execute('SELECT * FROM products WHERE id=?',(product_match.group(1),)).fetchone()
                if not target:return self.send_json({'error':'Produit introuvable.'},404)
                cur.execute('UPDATE products SET is_active=0,updated_at=? WHERE id=?',(now(),product_match.group(1)))
                cur.execute('INSERT INTO audit(username,action,details,created_at) VALUES(?,?,?,?)',
                    (u.get('username',''),'Archivage d’un produit',json.dumps({'table':'products','id':product_match.group(1),'before':dict(target)},ensure_ascii=False),now()))
            elif p.startswith('/api/users/') and p.rsplit('/',1)[-1].isdigit():
                if u['role'] not in ('Administrateur','Superadmin'):return self.send_json({'error':'Accès réservé à l’administrateur ou au Superadmin'},403)
                uid=int(p.rsplit('/',1)[-1])
                if uid==u['id']:return self.send_json({'error':'Vous ne pouvez pas supprimer votre propre compte'},400)
                target=c.execute('SELECT id,username,role FROM users WHERE id=?',(uid,)).fetchone()
                if not target:return self.send_json({'error':'Utilisateur introuvable'},404)
                if str(target['username']).lower()=='destrox' or target['role']=='Superadmin':
                    return self.send_json({'error':'Le compte Superadmin destrox est protégé et ne peut être supprimé par personne.'},403)
                if u['role']!='Superadmin' and target['role']=='Administrateur' and c.execute("SELECT COUNT(*) n FROM users WHERE role='Administrateur'").fetchone()['n']<=1:
                    return self.send_json({'error':'Impossible de supprimer le dernier administrateur'},400)
                c.execute('INSERT INTO audit(username,action,details,created_at) VALUES(?,?,?,?)',
                          (u.get('username',''),'Suppression d’un utilisateur',json.dumps({'table':'users','id':uid,'before':dict(target)},ensure_ascii=False),now()))
                c.execute('DELETE FROM users WHERE id=?',(uid,))
            elif p.startswith('/api/sales/') and p.rsplit('/',1)[-1].isdigit():
                uid=int(p.rsplit('/',1)[-1]); target=c.execute('SELECT * FROM sales WHERE id=?',(uid,)).fetchone()
                if not target:return self.send_json({'error':'Vente introuvable'},404)
                cur.execute('INSERT INTO audit(username,action,details,created_at) VALUES(?,?,?,?)',(u.get('username',''),'Suppression : vente',json.dumps({'table':'sales','id':uid,'before':dict(target)},ensure_ascii=False),now()))
                if target['product_id'] is not None:
                    q=max(1,int(target['quantity'] or 1)); cur.execute('UPDATE products SET stock_quantity=stock_quantity+?,updated_at=? WHERE id=?',(q,now(),target['product_id']))
                    cur.execute('INSERT INTO stock_movements(product_id,movement_type,quantity_delta,reference_table,reference_id,actor_id,actor_username,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(target['product_id'],'retour_vente',q,'sales',str(uid),u.get('id'),u.get('username',''),'Retour en stock après suppression de vente',now()))
                c.execute('DELETE FROM sales WHERE id=?',(uid,))
            elif p.startswith('/api/deposits/') and p.rsplit('/',1)[-1].isdigit():
                uid=int(p.rsplit('/',1)[-1]); target=c.execute('SELECT * FROM deposits WHERE id=?',(uid,)).fetchone()
                if not target:return self.send_json({'error':'Dépôt introuvable'},404)
                cur.execute('INSERT INTO audit(username,action,details,created_at) VALUES(?,?,?,?)',(u.get('username',''),'Suppression : dépôt',json.dumps({'table':'deposits','id':uid,'before':dict(target)},ensure_ascii=False),now()))
                c.execute('DELETE FROM deposits WHERE id=?',(uid,))
            elif p.startswith('/api/expenses/') and p.rsplit('/',1)[-1].isdigit():
                uid=int(p.rsplit('/',1)[-1]); target=c.execute('SELECT * FROM expenses WHERE id=?',(uid,)).fetchone()
                if not target:return self.send_json({'error':'Dépense introuvable'},404)
                cur.execute('INSERT INTO audit(username,action,details,created_at) VALUES(?,?,?,?)',(u.get('username',''),'Suppression : dépense',json.dumps({'table':'expenses','id':uid,'before':dict(target)},ensure_ascii=False),now()))
                c.execute('DELETE FROM expenses WHERE id=?',(uid,))
            else:return self.send_json({'error':'Introuvable'},404)
            c.commit(); return self.send_json({'ok':True})
        finally:c.close()

def cfg():
    try:
        c=json.load(open(CONFIG_PATH,encoding='utf-8'))
        c.setdefault('port',DEFAULT_PORT)
        changed=False
        # Old builds stored server_mode=True without an explicit network_mode.
        if 'network_mode' not in c:
            c['network_mode']='auto'; c['server_mode']=False; c['server']=''; changed=True
        # Repair the common copied-config problem: a local server config
        # (127.0.0.1 + server_mode=true) copied to every laptop. In AUTO mode
        # the application must discover the existing LAN server instead of
        # forcing every laptop to become a server.
        elif str(c.get('network_mode','')).lower()=='server' and str(c.get('server','')).strip() in ('','127.0.0.1','localhost') and bool(c.get('server_mode')):
            c['network_mode']='auto'; c['server_mode']=False; c['server']=''; changed=True
        if changed:
            try:
                with open(CONFIG_PATH,'w',encoding='utf-8') as f: json.dump(c,f,indent=2)
            except Exception: pass
        return c
    except Exception:
        return {'server':'','port':DEFAULT_PORT,'server_mode':False,'network_mode':'auto'}

def start_server(port):
    init_db(); backup_database(); start_backup_scheduler(); discovery_responder(port)
    s=ThreadingHTTPServer(('0.0.0.0',port),API); threading.Thread(target=s.serve_forever,daemon=True).start(); return s

def server_is_alive(host,port,timeout=0.8):
    try:
        with socket.create_connection((host,port),timeout=timeout): return True
    except OSError: return False

def start_local_cloud_api():
    # Every laptop serves its UI/API only on loopback; company records remain in Supabase.
    server=ThreadingHTTPServer(('127.0.0.1',0),API)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    return server

def show_window(url):
    try:
        import webview
        webview.create_window(APP_NAME,url,width=1500,height=900,min_size=(1100,700),resizable=True)
        webview.start()
    except ImportError:
        print('pywebview is not installed. Run: pip install -r requirements.txt')
        webbrowser.open(url)
        input('Press Enter to exit...')

def main():
    global RUNTIME_ROLE, RUNTIME_SERVER, RUNTIME_PORT
    # Cloud mode is opt-in per laptop via %USERPROFILE%\\WA_Manager_System\\cloud_config.json.
    # Without that file, the previous LAN/SQLite mode remains available as a fallback.
    cloud_cfg=load_cloud_config()
    if os.path.exists(CLOUD_CONFIG_PATH):
        if cloud_cfg.get('error'):
            message='Configuration Supabase invalide : '+str(cloud_cfg['error'])
            try:
                import ctypes
                ctypes.windll.user32.MessageBoxW(0,message,APP_NAME,0x10)
            except Exception:
                print(message)
            raise SystemExit(message)
        if cloud_cfg.get('enabled'):
            server=start_local_cloud_api()
            RUNTIME_ROLE='cloud'; RUNTIME_SERVER='Supabase'; RUNTIME_PORT=443
            url=f'http://127.0.0.1:{server.server_port}/'
            show_window(url)
            return
    c=cfg(); port=int(c.get('port',DEFAULT_PORT)); mode=str(c.get('network_mode','auto')).lower()
    configured_server=str(c.get('server','')).strip()
    server_mode=False; server_host=''

    if mode=='server' or bool(c.get('server_mode')):
        server_mode=True
    elif mode=='client':
        if configured_server and server_is_alive(configured_server,port):
            server_host=configured_server
        else:
            # If the configured server is unavailable, try discovery before failing over.
            found=discover_server()
            if found: server_host,port=found
            else: server_host=configured_server
    else:
        # Auto mode: find an existing server first. If none exists, wait a short randomized
        # period so two PCs started together don't both immediately become servers.
        found=discover_server()
        if found:
            server_host,port=found; save_cfg(server_host,port,'client')
        else:
            time.sleep(0.35 + (secrets.randbelow(1500)/1000.0))
            found=discover_server(timeout=1.0)
            if found:
                server_host,port=found; save_cfg(server_host,port,'client')
            else:
                server_mode=True; save_cfg(local_ip(),port,'server')

    if server_mode:
        start_server(port); RUNTIME_ROLE='server'; RUNTIME_SERVER=local_ip(); RUNTIME_PORT=port; url=f'http://127.0.0.1:{port}/'
    else:
        if not server_host:
            # Client mode should never silently create a second database.
            # If the configured server cannot be reached, retry discovery.
            found=discover_server(timeout=2.5)
            if found: server_host,port=found
        if not server_host:
            raise SystemExit('W-A Systeme de Gestion Server was not found on the LAN. Start the main computer first.')
        RUNTIME_ROLE='client'; RUNTIME_SERVER=server_host; RUNTIME_PORT=port; url=f'http://{server_host}:{port}/'

    show_window(url)
if __name__=='__main__':main()
